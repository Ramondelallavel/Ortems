"""Import and export of *every* table, with all its columns (Excel / CSV / JSON / database rows).

The friendly import templates in :mod:`imports` cover the common loads (items, routings, orders…).
This module adds a generic template for each entity of the master-data registry and for each of its
child tables (calendar shifts, BOM lines, order operations…), derived from the data model itself:

* one column per field; references are written as the referenced record's code (or number, name…)
  and also accept the record id;
* an ``id`` column, so an exported sheet can be edited and imported back (round trip);
* a ``delete`` column: ``yes`` deletes the row (records still referenced are deactivated instead,
  exactly as in the editor);
* rows are matched by ``id`` or by the table's natural key (its unique columns), else created.

Every row is written through the same service functions as the editor (:mod:`masterdata`), so the
same validation hooks, optimistic rules, audit trail and tenant isolation apply. Validation is a real
dry run: all rows are applied inside a savepoint that is rolled back, so the report shows exactly the
errors the import would hit.
"""

from __future__ import annotations

import json
import uuid
from dataclasses import dataclass, field
from datetime import UTC, date, datetime, time
from typing import Any
from zoneinfo import ZoneInfo

from sqlalchemy import JSON, Boolean, Date, DateTime, Float, Integer, Numeric, Time, UniqueConstraint, func, select
from sqlalchemy.orm import Session

from .. import models as M
from ..core.db import Base
from ..core.errors import DomainError, NotFound, ValidationFailed
from . import audit, masterdata
from .context import Ctx
from .imports import PREVIEW_ROWS, ConvError, F, Template, _to_kv, convert

PREFIX = "table:"
SYSTEM = masterdata.SYSTEM | {"tenant_id"}
DELETE_ALIASES = ("delete", "borrar", "eliminar", "_delete", "remove", "supprimer")
# column names commonly found in ERP/MES tables and spreadsheets (Spanish and English)
COMMON_ALIASES: dict[str, tuple[str, ...]] = {
    "code": ("codigo", "cod", "referencia", "ref", "clave", "sku"),
    "name": ("nombre", "descripcion", "description", "denominacion", "designation"),
    "number": ("numero", "num", "no", "nro", "order_number", "numero_orden", "pedido"),
    "description": ("descripcion", "desc", "detalle"),
    "priority": ("prioridad", "prio"),
    "quantity": ("cantidad", "cant", "qty", "unidades"),
    "item": ("articulo", "producto", "material", "item_code", "product", "codigo_articulo", "referencia_articulo"),
    "component": ("componente", "component_code", "codigo_componente"),
    "customer": ("cliente", "customer_code", "codigo_cliente"),
    "supplier": ("proveedor", "supplier_code", "codigo_proveedor"),
    "resource": ("recurso", "maquina", "machine", "resource_code", "codigo_recurso"),
    "machine": ("maquina", "machine_code"),
    "tool": ("herramienta", "utillaje", "tool_code"),
    "calendar": ("calendario", "calendar_code"),
    "plant": ("planta", "fabrica", "centro", "plant_code"),
    "area": ("area", "seccion", "zona"),
    "family": ("familia", "family_code"),
    "order": ("orden", "orden_fabricacion", "of", "order_number"),
    "operator": ("operario", "operator_code"),
    "due_date": ("fecha_entrega", "fecha_requerida", "fecha_compromiso", "due", "entrega"),
    "release_date": ("fecha_lanzamiento", "fecha_liberacion", "fecha_inicio_minima"),
    "start": ("inicio", "fecha_inicio", "desde", "from"),
    "end": ("fin", "fecha_fin", "hasta", "to"),
    "status": ("estado", "situacion"),
    "location": ("ubicacion", "almacen", "warehouse"),
    "on_hand": ("stock", "existencias", "disponible", "cantidad_stock"),
    "safety_stock": ("stock_seguridad", "stock_minimo"),
    "uom": ("unidad", "unidad_medida", "um"),
    "lead_time_days": ("plazo", "plazo_entrega", "lead_time", "dias_entrega"),
    "unit_cost": ("coste", "coste_unitario", "costo", "cost"),
    "unit_price": ("precio", "precio_unitario", "price"),
    "capacity": ("capacidad",),
    "efficiency": ("eficiencia", "rendimiento"),
    "country": ("pais",),
    "notes": ("notas", "observaciones", "comentarios"),
    "reason": ("motivo", "causa"),
    "kind": ("tipo", "clase"),
    "item_type": ("tipo_articulo", "tipo"),
    "seq": ("secuencia", "operacion", "fase", "op"),
    "line_no": ("linea", "num_linea", "line"),
    "expected_date": ("fecha_prevista", "fecha_recepcion", "fecha_llegada"),
    "period_start": ("periodo", "semana", "fecha", "period"),
    "good_quantity": ("cantidad_buena", "buenas", "good"),
    "scrap_quantity": ("rechazo", "chatarra", "scrap", "malas"),
    "setup_minutes": ("preparacion", "tiempo_preparacion", "setup"),
    "run_minutes_per_unit": ("tiempo_unitario", "minutos_unidad", "tiempo_ciclo", "run_time"),
}
YES = {"1", "true", "yes", "y", "si", "sí", "x", "delete", "borrar", "eliminar"}


@dataclass
class TableSpec:
    name: str  # "table:<entity>" or "table:<entity>.<child>"
    label: str
    group: str
    model: type
    edef: masterdata.EntityDef | None  # top-level entity (written through masterdata)
    parent: masterdata.EntityDef | None = None  # child table: its parent entity
    child: masterdata.ChildDef | None = None
    columns: list[Any] = field(default_factory=list)  # writable SQLAlchemy columns
    refs: dict[str, type] = field(default_factory=dict)  # column -> referenced model
    key: list[str] = field(default_factory=list)  # natural key columns
    template: Template | None = None
    read_perm: str = "masterdata:read"
    write_perm: str = "masterdata:write"

    @property
    def entity(self) -> str:
        return self.name[len(PREFIX) :]

    @property
    def sheet(self) -> str:
        return self.entity[:31]


def field_name(col: str) -> str:
    return col[:-3] if col.endswith("_id") else col


def _ftype(col: Any) -> str:
    t = col.type
    if isinstance(t, Boolean):
        return "bool"
    if isinstance(t, Integer):
        return "int"
    if isinstance(t, Float | Numeric):
        return "float"
    if isinstance(t, DateTime):
        return "datetime"
    if isinstance(t, Date):
        return "date"
    if isinstance(t, Time):
        return "time"
    if isinstance(t, JSON) or t.__class__.__name__ in ("JSON", "JSONB", "Variant"):
        return "json"
    return "str"


def _required(col: Any) -> bool:
    return not col.nullable and col.default is None and col.server_default is None and not col.primary_key


_MODELS_BY_TABLE: dict[str, type] = {}


def _model_for_table(name: str) -> type | None:
    if not _MODELS_BY_TABLE:
        for mapper in Base.registry.mappers:
            cls = mapper.class_
            if hasattr(cls, "__table__"):
                _MODELS_BY_TABLE.setdefault(cls.__table__.name, cls)
    return _MODELS_BY_TABLE.get(name)


def _build(name: str, label: str, group: str, model: type, edef, parent=None, child=None) -> TableSpec:
    spec = TableSpec(name=name, label=label, group=group, model=model, edef=edef, parent=parent, child=child)
    base = edef or parent
    spec.read_perm, spec.write_perm = base.read_perm, base.write_perm
    for col in model.__table__.columns:
        if col.key in SYSTEM:
            continue
        spec.columns.append(col)
        if col.foreign_keys:
            target = _model_for_table(next(iter(col.foreign_keys)).column.table.name)
            if target is not None:
                spec.refs[col.key] = target
    for con in model.__table__.constraints:
        if isinstance(con, UniqueConstraint):
            cols = [c.name for c in con.columns if c.name != "tenant_id"]
            if cols:
                spec.key = cols
                break
    fields = [F("id", description="Record id (keep it to update the same record; leave empty for new records)", aliases=("uuid", "record_id"))]
    for col in spec.columns:
        fname = field_name(col.key) if col.key in spec.refs else col.key
        req = _required(col)
        desc = ""
        if col.key in spec.refs:
            desc = f"{spec.refs[col.key].__tablename__.replace('_', ' ')}: code (or id)"
        if req:
            desc = (desc + "; " if desc else "") + "required for new records"
        if child is not None and col.key == child.fk:
            desc = f"parent {parent.label.lower()}: code (or id)"
        aliases = ((col.key,) if fname != col.key else ()) + COMMON_ALIASES.get(fname, ())
        fields.append(F(fname, "ref" if col.key in spec.refs else _ftype(col), required=False, description=desc, aliases=aliases))
    fields.append(F("delete", "bool", description="yes = delete this record", aliases=DELETE_ALIASES))
    key_fields = tuple(field_name(k) if k in spec.refs else k for k in spec.key) or ("id",)
    spec.template = Template(name, label, fields, key_fields, f"All columns of {label.lower()}. Rows are matched by id or by {', '.join(key_fields)}; the rest are created.", read_perm=spec.read_perm, write_perm=spec.write_perm, unique_rows=False)
    return spec


_SPECS: dict[str, TableSpec] = {}


def specs() -> dict[str, TableSpec]:
    if not _SPECS:
        for d in masterdata.REGISTRY.values():
            _SPECS[PREFIX + d.name] = _build(PREFIX + d.name, d.label, d.group, d.model, d)
            for ch in d.children:
                n = f"{PREFIX}{d.name}.{ch.name}"
                _SPECS[n] = _build(n, f"{d.label} — {ch.name.replace('_', ' ')}", d.group, ch.model, None, parent=d, child=ch)
    return _SPECS


def get_spec(name: str) -> TableSpec:
    if not name.startswith(PREFIX):
        name = PREFIX + name
    sp = specs().get(name)
    if sp is None:
        raise NotFound(f"Unknown table '{name[len(PREFIX):]}'", code="UNKNOWN_ENTITY")
    return sp


def is_table(entity: str) -> bool:
    return entity.startswith(PREFIX)


# canonical tables for a whole-database workbook (no filtered aliases such as machines/products)
def workbook_tables() -> list[TableSpec]:
    out = []
    for sp in specs().values():
        base = sp.edef or sp.parent
        if base.fixed_filter:
            continue
        out.append(sp)
    return sorted(out, key=dependency_rank)


def dependency_rank(sp: TableSpec) -> int:
    order = {t.name: i for i, t in enumerate(Base.metadata.sorted_tables)}
    return order.get(sp.model.__table__.name, 10_000)


# =============================================================================================
# reference lookups (by code / number / name, or id)
# =============================================================================================


def _label_rows(s: Session, model: type) -> list[tuple[str, uuid.UUID, Any]]:
    """(label, id, plant_id) for a referenced model."""
    pid = getattr(model, "plant_id", None)
    if model in (M.Routing, M.Bom):
        cols = [M.Item.code, model.id] + ([pid] if pid is not None else [])
        q = select(*cols).join(M.Item, M.Item.id == model.item_id).where(model.is_active.is_(True))
        return [(str(r[0]), r[1], r[2] if len(r) > 2 else None) for r in s.execute(q)]
    for c in ("code", "number", "username", "lot_number", "name"):
        if hasattr(model, c):
            q = select(getattr(model, c), model.id, pid) if pid is not None else select(getattr(model, c), model.id)
            return [(str(r[0]), r[1], r[2] if len(r) > 2 else None) for r in s.execute(q) if r[0] is not None]
    return []


class RefIndex:
    def __init__(self, s: Session, plant: M.Plant):
        self.s, self.plant = s, plant
        self._by_label: dict[type, dict[str, list[tuple[uuid.UUID, Any]]]] = {}
        self._label_of: dict[type, dict[uuid.UUID, str]] = {}
        self._ids: dict[type, set[uuid.UUID]] = {}

    def _load(self, model: type) -> None:
        if model in self._by_label:
            return
        by: dict[str, list[tuple[uuid.UUID, Any]]] = {}
        rows = _label_rows(self.s, model)
        for label, id_, pid in rows:
            by.setdefault(label.lower(), []).append((id_, pid))
        self._by_label[model] = by
        self._label_of[model] = {}
        for label, id_, _pid in rows:
            if len(by[label.lower()]) == 1:
                self._label_of[model][id_] = label

    def resolve(self, model: type, value: Any) -> uuid.UUID:
        text = str(value).strip()
        try:
            rid = uuid.UUID(text)
        except ValueError:
            rid = None
        if rid is not None:
            if model not in self._ids:
                self._ids[model] = set(self.s.scalars(select(model.id)))
            if rid not in self._ids[model]:
                raise ConvError(f"no {model.__tablename__.replace('_', ' ')} with id {text}")
            return rid
        self._load(model)
        cands = self._by_label[model].get(text.lower(), [])
        if len(cands) > 1:
            here = [c for c in cands if c[1] in (None, self.plant.id)]
            cands = here if here else cands
        if not cands:
            raise ConvError(f"unknown {model.__tablename__.replace('_', ' ')} '{text}'")
        if len(cands) > 1:
            raise ConvError(f"'{text}' matches {len(cands)} {model.__tablename__.replace('_', ' ')} records; use the id")
        return cands[0][0]

    def label(self, model: type, id_: uuid.UUID | None) -> str | None:
        if id_ is None:
            return None
        self._load(model)
        return self._label_of[model].get(id_) or str(id_)

    def forget(self, model: type) -> None:
        self._by_label.pop(model, None)
        self._label_of.pop(model, None)
        self._ids.pop(model, None)


# =============================================================================================
# conversion
# =============================================================================================


def _convert_json(raw: Any) -> Any:
    if isinstance(raw, dict | list):
        return raw
    text = str(raw).strip()
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        return _to_kv(text)


def convert_rows(s: Session, job: M.ImportJob, spec: TableSpec, plant: M.Plant, refs: RefIndex) -> tuple[list[tuple[int, dict[str, Any]]], list[dict[str, Any]]]:
    """→ ([(file row, {column|'id'|'delete': value})], errors)."""
    tz = ZoneInfo(plant.timezone)
    date_format = (job.options or {}).get("date_format", "DMY")
    col_idx = {c: i for i, c in enumerate(job.columns or [])}
    fmap = {f.name: col_idx.get(job.mapping.get(f.name)) if job.mapping.get(f.name) else None for f in spec.template.fields}
    by_field = {(field_name(c.key) if c.key in spec.refs else c.key): c for c in spec.columns}
    out, errors = [], []
    for n, raw in enumerate(job.rows or [], start=2):
        vals: dict[str, Any] = {}
        row_errs = []
        for f in spec.template.fields:
            idx = fmap.get(f.name)
            if idx is None:
                continue
            rv = raw[idx] if idx < len(raw) else None
            empty = rv is None or (isinstance(rv, str) and not rv.strip())
            col = by_field.get(f.name)
            try:
                if f.name == "id":
                    if not empty:
                        vals["id"] = uuid.UUID(str(rv).strip())
                    continue
                if f.name == "delete":
                    vals["delete"] = (not empty) and str(rv).strip().lower() in YES
                    continue
                if empty:
                    vals[col.key] = None
                elif col.key in spec.refs:
                    vals[col.key] = refs.resolve(spec.refs[col.key], rv)
                elif f.type == "json":
                    vals[col.key] = _convert_json(rv)
                else:
                    vals[col.key] = convert(f, rv, tz, date_format)
            except (ConvError, ValueError) as exc:
                row_errs.append({"row": n, "field": f.name, "value": rv, "message": str(exc) or "invalid value"})
        if row_errs:
            errors.extend(row_errs)
        else:
            out.append((n, vals))
    return out, errors


# =============================================================================================
# apply
# =============================================================================================


class Applier:
    def __init__(self, s: Session, ctx: Ctx, spec: TableSpec, plant: M.Plant, mode: str):
        self.s, self.ctx, self.spec, self.plant, self.mode = s, ctx, spec, plant, mode
        self.loaded: dict[uuid.UUID, Any] = {}
        self.counts = {"created": 0, "updated": 0, "unchanged": 0, "deleted": 0, "deactivated": 0, "skipped": 0}

    def _find(self, vals: dict[str, Any]) -> Any:
        sp, s = self.spec, self.s
        if vals.get("id"):
            obj = self.loaded.get(vals["id"]) or s.get(sp.model, vals["id"])
            if obj is None:
                raise ValidationFailed(f"No {sp.label.lower()} with id {vals['id']}", code="NOT_FOUND", context={"field": "id"})
            return obj
        if sp.key and all(vals.get(k) is not None for k in sp.key):
            q = select(sp.model).where(*[getattr(sp.model, k) == vals[k] for k in sp.key])
            if hasattr(sp.model, "code") and "code" in sp.key and isinstance(vals.get("code"), str):
                q = select(sp.model).where(*[(func.lower(sp.model.code) == vals["code"].lower()) if k == "code" else getattr(sp.model, k) == vals[k] for k in sp.key])
            return s.scalar(q)
        return None

    def quick(self, vals: dict[str, Any]) -> str | None:
        """'unchanged' / 'skipped' when the row needs no write (no savepoint, no audit), else None."""
        if vals.get("delete"):
            return None
        obj = self._find(vals)
        if obj is not None and self.mode == "CREATE_ONLY" or obj is None and self.mode == "UPDATE_ONLY":
            return "skipped"
        if obj is not None and not _changes(obj, self.spec, {k: v for k, v in vals.items() if k not in ("id", "delete")}):
            return "unchanged"
        return None

    def apply(self, vals: dict[str, Any]) -> str:
        sp, s, ctx = self.spec, self.s, self.ctx
        obj = self._find(vals)
        creating = obj is None
        data = {k: v for k, v in vals.items() if k not in ("id", "delete")}
        if vals.get("delete"):
            if obj is None:
                raise ValidationFailed("Row marked for deletion does not exist", code="NOT_FOUND")
            if sp.edef is not None:
                out = masterdata.delete_row(s, ctx, sp.edef.name, obj.id, reason="import")
                key = "deactivated" if out.get("deactivated") else "deleted"
            else:
                ctx.require(sp.write_perm)
                before = audit.snapshot(obj)
                s.delete(obj)
                s.flush()
                audit.record(s, ctx, "DELETE", sp.entity, obj.id, masterdata._label(obj), before=before, reason="import")
                key = "deleted"
            return key
        if obj is not None and self.mode == "CREATE_ONLY" or obj is None and self.mode == "UPDATE_ONLY":
            return "skipped"
        if obj is None:
            if "plant_id" in {c.key for c in sp.columns} and data.get("plant_id") is None:
                data["plant_id"] = self.plant.id
            for c in sp.columns:  # column defaults are applied now so validation hooks see them
                if data.get(c.key) is None and c.default is not None and getattr(c.default, "is_scalar", False):
                    data[c.key] = c.default.arg
            missing = [field_name(c.key) if c.key in sp.refs else c.key for c in sp.columns if _required(c) and data.get(c.key) is None]
            if missing:
                raise ValidationFailed(f"Required for new records: {', '.join(missing)}", code="REQUIRED", context={"field": missing[0]})
        # an empty cell never clears a mandatory column
        data = {k: v for k, v in data.items() if v is not None or next(c for c in sp.columns if c.key == k).nullable}
        if obj is not None and not _changes(obj, sp, data):
            return "unchanged"
        if sp.edef is not None:
            if obj is None:
                masterdata.create_row(s, ctx, sp.edef.name, data, reason="import", return_row=False)
            else:
                masterdata.update_row(s, ctx, sp.edef.name, obj.id, data, reason="import", return_row=False)
        else:
            ctx.require(sp.write_perm)
            if creating:
                obj = sp.model()
                obj.tenant_id = ctx.tenant_id
                s.add(obj)
            before = None if creating else audit.snapshot(obj)
            masterdata.apply_fields(obj, data, None, creating=True)
            s.flush()
            audit.record(s, ctx, "CREATE" if creating else "UPDATE", sp.entity, obj.id, masterdata._label(obj), before=before, after=audit.snapshot(obj), reason="import")
        return "created" if creating else "updated"


def _norm(v: Any) -> Any:
    if isinstance(v, datetime):
        return (v.astimezone(UTC) if v.tzinfo else v.replace(tzinfo=UTC)).replace(tzinfo=None, microsecond=0)
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def _changes(obj: Any, spec: TableSpec, data: dict[str, Any]) -> bool:
    """Would writing ``data`` change the record? (unchanged rows are not rewritten nor audited)"""
    cols = {c.key: c for c in spec.columns}
    for k, v in data.items():
        col = cols.get(k)
        if col is None:
            continue
        try:
            nv = masterdata._coerce(col, v)
        except DomainError:
            return True
        if _norm(getattr(obj, k)) != _norm(nv):
            return True
    return False


def _error_of(exc: Exception) -> tuple[str | None, str]:
    if isinstance(exc, DomainError):
        fld = (exc.context or {}).get("field")
        return (field_name(fld) if isinstance(fld, str) else None), exc.message
    return None, f"{exc.__class__.__name__}: {exc}"


def run(s: Session, ctx: Ctx, job: M.ImportJob, plant: M.Plant, refs: RefIndex | None = None) -> dict[str, Any]:
    """Apply the job's rows (callers wrap this in a savepoint and decide to keep or roll back).
    Returns counts, per-row errors and a preview of the actions."""
    spec = get_spec(job.entity)
    ctx.require("integration:import")
    ctx.require(spec.write_perm)
    refs = refs or RefIndex(s, plant)
    rows, errors = convert_rows(s, job, spec, plant, refs)
    mode = (job.options or {}).get("mode", "UPSERT")
    ap = Applier(s, ctx, spec, plant, mode)
    preview = []
    ids = [v["id"] for _n, v in rows if v.get("id")]
    for i in range(0, len(ids), 500):  # load the records once, in batches (kept referenced while the rows are applied)
        for o in s.scalars(select(spec.model).where(spec.model.id.in_(ids[i : i + 500]))):
            ap.loaded[o.id] = o
    for n, vals in rows:
        try:
            key = ap.quick(vals)
            if key is None:
                with s.begin_nested():
                    key = ap.apply(vals)
            ap.counts[key] += 1
            action = key.upper()
        except Exception as exc:  # noqa: BLE001 - reported per row
            fld, msg = _error_of(exc)
            errors.append({"row": n, "field": fld, "value": None, "message": msg})
            action = "ERROR"
        if len(preview) < PREVIEW_ROWS:
            preview.append({"row": n, "action": action, "values": {k: _pv(v) for k, v in vals.items()}})
    refs.forget(spec.model)
    return {"counts": ap.counts, "errors": errors, "preview": preview}


def _pv(v: Any) -> Any:
    if isinstance(v, datetime | date | time):
        return v.isoformat()
    if isinstance(v, uuid.UUID):
        return str(v)
    return v


def stats_from(job: M.ImportJob, result: dict[str, Any]) -> dict[str, Any]:
    errors = result["errors"]
    c = result["counts"]
    error_rows = len({e["row"] for e in errors if e["row"] is not None})
    total = len(job.rows or [])
    blocking = bool(errors) and not (job.options or {}).get("skip_invalid_rows") or any(e["row"] is None for e in errors)
    good = total - error_rows
    return {
        "rows": total,
        "valid_rows": good,
        "error_rows": error_rows,
        "errors": len(errors),
        "warnings": 0,
        "to_create": c["created"],
        "to_update": c["updated"],
        "unchanged": c["unchanged"],
        "to_delete": c["deleted"] + c["deactivated"],
        "to_skip": c["skipped"],
        "can_import": not blocking and good > 0,
        "blocking_reason": ("Fix the errors or choose to import only the valid rows." if errors else None) if blocking else None,
        "preview": result["preview"],
    }


# =============================================================================================
# export
# =============================================================================================


def _cell_out(v: Any, tz: ZoneInfo) -> Any:
    if isinstance(v, datetime):  # stored in UTC (SQLite hands it back without tzinfo)
        return (v if v.tzinfo else v.replace(tzinfo=UTC)).astimezone(tz).replace(tzinfo=None)
    if isinstance(v, uuid.UUID):
        return str(v)
    if isinstance(v, dict | list):
        return json.dumps(v, ensure_ascii=False, default=str)
    return v


def export_rows(s: Session, ctx: Ctx, name: str, plant: M.Plant | None) -> tuple[list[str], list[list[Any]]]:
    spec = get_spec(name)
    ctx.require(spec.read_perm)
    tz = ZoneInfo(plant.timezone if plant else "UTC")
    refs = RefIndex(s, plant) if plant else None
    q = select(spec.model)
    for k, v in (spec.edef.fixed_filter if spec.edef else {}).items():
        q = q.where(getattr(spec.model, k).in_(v) if isinstance(v, list | tuple) else getattr(spec.model, k) == v)
    if spec.edef and spec.edef.plant_scoped and plant is not None and hasattr(spec.model, "plant_id"):
        q = q.where(spec.model.plant_id == plant.id)
    order = (spec.edef.order_by if spec.edef else (spec.child.order_by or None)) or None
    if spec.child is not None:
        q = q.order_by(getattr(spec.model, spec.child.fk))
    if order and hasattr(spec.model, order.lstrip("-")):
        col = getattr(spec.model, order.lstrip("-"))
        q = q.order_by(col.desc() if order.startswith("-") else col)
    cols = [f.name for f in spec.template.fields if f.name != "delete"]
    by_field = {(field_name(c.key) if c.key in spec.refs else c.key): c for c in spec.columns}
    rows = []
    for obj in s.scalars(q):
        r = []
        for fname in cols:
            if fname == "id":
                r.append(str(obj.id))
                continue
            col = by_field[fname]
            v = getattr(obj, col.key)
            if col.key in spec.refs and v is not None:
                r.append(refs.label(spec.refs[col.key], v) if refs else str(v))
            else:
                r.append(_cell_out(v, tz))
        rows.append(r)
    return cols, rows


def xlsx(sheets: list[tuple[str, list[str], list[list[Any]], list[F] | None]], title: str) -> bytes:
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    for name, cols, rows, fields in sheets:
        ws = wb.create_sheet(name[:31])
        ws.append(cols)
        for i, c in enumerate(ws[1], start=1):
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="12203A")
            ws.column_dimensions[get_column_letter(i)].width = max(10, min(40, len(str(c.value)) + 4))
        for r in rows:
            ws.append(r)
        ws.freeze_panes = "B2" if cols and cols[0] == "id" else "A2"
        if cols and cols[0] == "id":
            ws.column_dimensions["A"].hidden = False
            ws.column_dimensions["A"].width = 12
    help_ws = wb.create_sheet("How to use")
    for line in [
        [title],
        [],
        ["Edit the rows and import the file again (Integrations → Import, or the Import button of each section)."],
        ["Keep the id column to update the same records. New rows: leave id empty."],
        ["Add a column named delete and write yes to delete a row (records still in use are deactivated instead)."],
        ["References (item, resource, calendar…) are written as codes; ids are accepted too."],
        ["Dates and times are in the plant's local time."],
    ]:
        help_ws.append(line)
    help_ws.column_dimensions["A"].width = 110
    help_ws["A1"].font = Font(bold=True, size=13)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
