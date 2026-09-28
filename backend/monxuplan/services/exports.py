"""Plan exports: Excel workbook (Schedule, Orders, Operations, Capacity, Materials, Alerts, KPIs),
single-sheet CSV and JSON. Every value comes from the stored plan version — nothing is recomputed
differently from what the planner sees on screen. Times are written in the plant's local time
(Excel has no time zones); the time zone is stated in the header of every sheet.
"""

from __future__ import annotations

import csv
import io
import json
import uuid
from collections.abc import Iterable
from datetime import UTC, datetime
from typing import Any
from zoneinfo import ZoneInfo

from sqlalchemy import select
from sqlalchemy.orm import Session

from ..core.errors import ValidationFailed
from ..models import Alert, ConstraintViolation, Customer, ExportJob, Item, Plant, ProductionOrder, ProductionOrderOperation, Resource, ScheduledOperation
from .analytics import KPI_CATALOGUE
from .context import Ctx
from .engine_view import replay
from .views import _get_plan

SHEETS = ["Schedule", "Orders", "Operations", "Capacity", "Materials", "Alerts", "KPIs"]


class _Table:
    def __init__(self, title: str, columns: list[str], rows: Iterable[list[Any]]):
        self.title = title
        self.columns = columns
        self.rows = list(rows)


def _local(dt: datetime | str | None, tz: ZoneInfo) -> datetime | None:
    if dt is None:
        return None
    if isinstance(dt, str):
        dt = datetime.fromisoformat(dt.replace("Z", "+00:00"))
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=UTC)
    return dt.astimezone(tz).replace(tzinfo=None)


def _h(minutes: float | int | None) -> float | None:
    return None if minutes is None else round(minutes / 60.0, 2)


# ---------------------------------------------------------------------------------------------
# sheet builders
# ---------------------------------------------------------------------------------------------


def _schedule(s: Session, plan, tz: ZoneInfo, res_codes: dict[str, str]) -> _Table:
    rows = list(s.scalars(select(ScheduledOperation).where(ScheduledOperation.plan_id == plan.id).order_by(ScheduledOperation.resource_key, ScheduledOperation.setup_start)))
    orders = {o.id: o for o in s.scalars(select(ProductionOrder).where(ProductionOrder.plant_id == plan.plant_id))}
    items = {i.id: i for i in s.scalars(select(Item))}
    poos = {p.id: p for p in s.scalars(select(ProductionOrderOperation).where(ProductionOrderOperation.id.in_({r.order_operation_id for r in rows if r.order_operation_id})))} if rows else {}
    out = []
    for r in rows:
        o = orders.get(r.order_id)
        it = items.get(o.item_id) if o else None
        poo = poos.get(r.order_operation_id)
        sec = ", ".join(res_codes.get(x.get("resource_id"), x.get("resource_id", "")) for x in (r.secondary or []) if isinstance(x, dict))
        b = r.binding or {}
        out.append(
            [
                res_codes.get(r.resource_key, r.resource_key),
                r.op_key,
                o.number if o else r.order_key,
                it.code if it else None,
                poo.name if poo else None,
                r.quantity,
                _local(r.setup_start, tz),
                _local(r.start, tz),
                _local(r.end, tz),
                r.setup_minutes,
                r.run_minutes,
                r.overtime_minutes,
                sec or None,
                r.zone,
                "yes" if r.is_locked else "",
                r.fixed_reason or ("yes" if r.is_fixed else ""),
                "yes" if r.is_late else "",
                "yes" if r.subcontracted else "",
                b.get("type"),
                b.get("message"),
            ]
        )
    return _Table(
        "Schedule",
        ["Resource", "Operation ID", "Order", "Item", "Operation", "Quantity", "Setup start", "Start", "End", "Setup (min)", "Run (min)", "Overtime (min)", "Secondary resources", "Zone", "Locked", "Fixed", "Late", "Subcontracted", "Binding constraint", "Binding detail"],
        out,
    )


def _orders(s: Session, plan, tz: ZoneInfo) -> _Table:
    an = plan.analysis or {}
    late = {x["order_id"]: x for x in (plan.kpi_details or {}).get("late_orders", [])}
    db = {str(o.id): o for o in s.scalars(select(ProductionOrder).where(ProductionOrder.plant_id == plan.plant_id))}
    items = {i.id: i for i in s.scalars(select(Item))}
    custs = {c.id: c for c in s.scalars(select(Customer))}
    out = []
    for o in sorted(an.get("orders", []), key=lambda x: x.get("due") or ""):
        d = db.get(o["order_id"])
        it = items.get(d.item_id) if d else None
        li = late.get(o["order_id"], {})
        out.append(
            [
                o.get("number"),
                it.code if it else None,
                it.name if it else None,
                d.quantity if d else None,
                custs[d.customer_id].name if d and d.customer_id in custs else None,
                d.priority if d else None,
                "yes" if d and d.expedite else "",
                _local(o.get("due"), tz),
                _local(o.get("start"), tz),
                _local(o.get("end"), tz),
                o.get("status"),
                _h(o.get("lateness_minutes")),
                o.get("material_status"),
                (li.get("cause") or {}).get("category"),
                (li.get("cause") or {}).get("text"),
            ]
        )
    return _Table("Orders", ["Order", "Item", "Description", "Quantity", "Customer", "Priority", "Expedite", "Due", "Planned start", "Planned end", "Status", "Lateness (h)", "Material status", "Delay cause", "Explanation"], out)


def _operations(s: Session, plan, tz: ZoneInfo, res_codes: dict[str, str]) -> _Table:
    sched = _schedule(s, plan, tz, res_codes)
    out = [[r[1], r[2], r[3], r[4], r[0], r[5], r[7], r[8], "SCHEDULED", None] for r in sched.rows]
    numbers = {str(i): n for i, n in s.execute(select(ProductionOrder.id, ProductionOrder.number).where(ProductionOrder.plant_id == plan.plant_id))}
    for u in (plan.analysis or {}).get("unscheduled", []):
        out.append([u.get("op_id"), numbers.get(u.get("order_id"), u.get("order_id")), None, None, None, None, None, None, "UNSCHEDULED", f"{u.get('reason')}: {u.get('message', '')}"])
    return _Table("Operations", ["Operation ID", "Order", "Item", "Operation", "Resource", "Quantity", "Start", "End", "State", "Reason"], out)


def _capacity(s: Session, plan, tz: ZoneInfo) -> _Table:
    from monxuplan_engine.capacity import load_profile

    cp, res = replay(s, plan)
    bks, loads = load_profile(cp, res.placements, res.timing, "day", None, None)
    out = []
    for rl in loads:
        r = cp.resources[rl.resource]
        if r.kind in ("SUBCONTRACTOR", "STORAGE", "TRANSPORT"):
            continue
        for bk, b in zip(bks, rl.buckets, strict=True):
            if not b["capacity"] and not b["scheduled"] and not b["requirement"]:
                continue
            out.append(
                [
                    r.code,
                    r.kind,
                    r.area,
                    _local(cp.dt(bk.start), tz).date(),
                    _h(b["capacity"]),
                    _h(b["scheduled"]),
                    _h(b["requirement"]),
                    _h(b.get("available")),
                    _h(b.get("overload")),
                    round(100 * b["utilization"], 1) if b.get("utilization") is not None else None,
                    b.get("state"),
                ]
            )
    return _Table("Capacity", ["Resource", "Kind", "Area", "Day", "Capacity (h)", "Scheduled (h)", "Requirement (h)", "Available (h)", "Overload (h)", "Utilisation (%)", "State"], out)


def _materials(s: Session, plan, tz: ZoneInfo) -> _Table:
    items = {str(i.id): i for i in s.scalars(select(Item))}
    orders = {str(i): n for i, n in s.execute(select(ProductionOrder.id, ProductionOrder.number).where(ProductionOrder.plant_id == plan.plant_id))}
    out = []
    for p in (plan.analysis or {}).get("pegging", []):
        it = items.get(p["material_id"])
        late = bool(p.get("supply_time") and p.get("need_time") and p["supply_time"] > p["need_time"])
        out.append(
            [
                orders.get(p["consumer_order_id"], p["consumer_order_id"]),
                p.get("consumer_op_id"),
                it.code if it else p["material_id"],
                it.uom if it else None,
                round(float(p["quantity"]), 4),
                p.get("supply_kind"),
                p.get("supply_ref") or orders.get(p.get("supply_order_id") or ""),
                _local(p.get("supply_time"), tz),
                _local(p.get("need_time"), tz),
                "yes" if late else "",
            ]
        )
    for u in (plan.analysis or {}).get("unscheduled", []):
        if u.get("reason") != "MATERIAL_SHORTAGE":
            continue
        for m in (u.get("details") or {}).get("materials", []):
            out.append([orders.get(u.get("order_id"), u.get("order_id")), u.get("op_id"), m.get("material"), m.get("uom"), m.get("required"), "SHORTAGE", f"shortfall {m.get('shortfall')}", None, None, "shortage"])
    return _Table("Materials", ["Consumer order", "Operation", "Material", "UoM", "Quantity", "Supply kind", "Supply reference", "Available at", "Needed at", "Late / shortage"], out)


def _alerts(s: Session, plan, tz: ZoneInfo) -> _Table:
    out = []
    for v in s.scalars(select(ConstraintViolation).where(ConstraintViolation.plan_id == plan.id).order_by(ConstraintViolation.severity, ConstraintViolation.type)):
        out.append(["VIOLATION", v.severity, v.hardness, v.type, v.message, v.order_key, v.op_key, v.resource_key, _local(v.start, tz), _local(v.end, tz)])
    for a in s.scalars(select(Alert).where(Alert.plan_id == plan.id).order_by(Alert.severity)):
        out.append(["ALERT", a.severity, None, a.type, f"{a.title} — {a.message}", None, None, None, _local(a.created_at, tz), None])
    return _Table("Alerts", ["Source", "Severity", "Hardness", "Type", "Message", "Order", "Operation", "Resource", "From", "To"], out)


def _kpis(plan) -> _Table:
    values = plan.kpis or {}
    out = [[k["group"], k["label"], values.get(k["code"]), k["unit"], k["code"]] for k in KPI_CATALOGUE if k["code"] in values]
    md = plan.solver_metadata or {}
    out += [
        ["Solver", "Provider", md.get("provider"), "", "provider"],
        ["Solver", "Status", md.get("status"), "", "status"],
        ["Solver", "Objective", md.get("objective"), "", "objective"],
        ["Solver", "Gap", md.get("gap"), "", "gap"],
        ["Solver", "Runtime", md.get("runtime_s"), "s", "runtime_s"],
        ["Solver", "Input hash", md.get("input_hash"), "", "input_hash"],
    ]
    return _Table("KPIs", ["Group", "Indicator", "Value", "Unit", "Code"], out)


def build_tables(s: Session, plan, sheets: list[str]) -> list[_Table]:
    plant = s.get(Plant, plan.plant_id)
    tz = ZoneInfo(plant.timezone if plant else "UTC")
    res_codes = {str(r.id): r.code for r in s.scalars(select(Resource))}
    out = []
    for name in sheets:
        if name == "Schedule":
            out.append(_schedule(s, plan, tz, res_codes))
        elif name == "Orders":
            out.append(_orders(s, plan, tz))
        elif name == "Operations":
            out.append(_operations(s, plan, tz, res_codes))
        elif name == "Capacity":
            out.append(_capacity(s, plan, tz))
        elif name == "Materials":
            out.append(_materials(s, plan, tz))
        elif name == "Alerts":
            out.append(_alerts(s, plan, tz))
        elif name == "KPIs":
            out.append(_kpis(plan))
    return out


# ---------------------------------------------------------------------------------------------
# writers
# ---------------------------------------------------------------------------------------------


def _xlsx(tables: list[_Table], title: str, tz_name: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    head_fill = PatternFill("solid", fgColor="12203A")
    head_font = Font(bold=True, color="FFFFFF")
    for t in tables:
        ws = wb.create_sheet(t.title)
        ws.append([f"{title} — {t.title}", f"Times in {tz_name}"])
        ws["A1"].font = Font(bold=True, size=12)
        ws.append(t.columns)
        for c in range(1, len(t.columns) + 1):
            cell = ws.cell(row=2, column=c)
            cell.fill = head_fill
            cell.font = head_font
            cell.alignment = Alignment(vertical="center")
        widths = [len(c) for c in t.columns]
        for row in t.rows:
            ws.append(row)
            for i, v in enumerate(row):
                n = 16 if isinstance(v, datetime) else len(str(v)) if v is not None else 0
                if n > widths[i]:
                    widths[i] = n
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = min(60, max(8, w + 2))
        for row in ws.iter_rows(min_row=3):
            for cell in row:
                if isinstance(cell.value, datetime):
                    cell.number_format = "yyyy-mm-dd hh:mm"
        ws.freeze_panes = "A3"
        if t.rows:
            ws.auto_filter.ref = f"A2:{get_column_letter(len(t.columns))}{len(t.rows) + 2}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _csv(t: _Table) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(t.columns)
    for r in t.rows:
        w.writerow(["" if v is None else v.isoformat(sep=" ", timespec="minutes") if isinstance(v, datetime) else v for v in r])
    return ("﻿" + buf.getvalue()).encode("utf-8")


def _jsonable(v: Any) -> Any:
    if isinstance(v, datetime):
        return v.isoformat(timespec="minutes")
    if hasattr(v, "isoformat"):
        return v.isoformat()
    return v


def export_plan(s: Session, ctx: Ctx, plan_id: uuid.UUID, fmt: str = "xlsx", sheet: str | None = None) -> tuple[bytes, str, str]:
    """Returns (bytes, media type, filename)."""
    ctx.require("integration:export")
    plan = _get_plan(s, ctx, plan_id)
    plant = s.get(Plant, plan.plant_id)
    tz_name = plant.timezone if plant else "UTC"
    if fmt == "xlsx":
        sheets = SHEETS if not sheet or sheet == "all" else [sheet]
    else:
        sheets = [sheet or "Schedule"]
    unknown = [x for x in sheets if x not in SHEETS]
    if unknown:
        raise ValidationFailed(f"Unknown sheet {unknown[0]!r}. Available: {', '.join(SHEETS)}", code="UNKNOWN_SHEET")
    tables = build_tables(s, plan, sheets)
    base = f"{plan.number}"
    if fmt == "xlsx":
        data, media, fname = _xlsx(tables, f"MonxuPlan {plan.number}", tz_name), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", f"{base}.xlsx"
    elif fmt == "csv":
        data, media, fname = _csv(tables[0]), "text/csv; charset=utf-8", f"{base}-{tables[0].title.lower()}.csv"
    elif fmt == "json":
        payload = {"plan": plan.number, "timezone": tz_name, "sheet": tables[0].title, "columns": tables[0].columns, "rows": [[_jsonable(v) for v in r] for r in tables[0].rows]}
        data, media, fname = json.dumps(payload, ensure_ascii=False).encode(), "application/json", f"{base}-{tables[0].title.lower()}.json"
    else:
        raise ValidationFailed("format must be xlsx, csv or json")
    s.add(ExportJob(tenant_id=ctx.tenant_id, kind="PLAN", file_format=fmt, params={"plan_id": str(plan.id), "sheets": sheets}, rows=sum(len(t.rows) for t in tables), target=fname))
    return data, media, fname


def export_entity(s: Session, ctx: Ctx, entity: str, fmt: str = "csv", plant_id: uuid.UUID | None = None) -> tuple[bytes, str, str]:
    """Master/transactional data export (same columns as the import templates)."""
    from .imports import TEMPLATES, template_rows

    if entity not in TEMPLATES:
        raise ValidationFailed(f"No export available for '{entity}'. Available: {', '.join(TEMPLATES)}", code="UNKNOWN_ENTITY")
    tpl = TEMPLATES[entity]
    ctx.require(tpl.read_perm)
    if plant_id:
        ctx.require_plant(plant_id)
    plant = s.get(Plant, plant_id) if plant_id else s.scalar(select(Plant).order_by(Plant.code).limit(1))
    tz_name = plant.timezone if plant else "UTC"
    tz = ZoneInfo(tz_name)
    cols = [f.name for f in tpl.fields]
    rows = [{k: (_local(v, tz) if isinstance(v, datetime) else v) for k, v in r.items()} for r in template_rows(s, entity, plant_id)]
    t = _Table(entity, cols, ([r.get(c) for c in cols] for r in rows))
    if fmt == "xlsx":
        data, media, fname = _xlsx([t], f"MonxuPlan {entity}", tz_name), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", f"{entity}.xlsx"
    elif fmt == "json":
        data, media, fname = json.dumps([{c: _jsonable(r.get(c)) for c in cols} for r in rows], ensure_ascii=False).encode(), "application/json", f"{entity}.json"
    else:
        data, media, fname = _csv(t), "text/csv; charset=utf-8", f"{entity}.csv"
    s.add(ExportJob(tenant_id=ctx.tenant_id, kind=entity.upper(), file_format=fmt, params={"plant_id": str(plant_id) if plant_id else None}, rows=len(t.rows), target=fname))
    return data, media, fname
