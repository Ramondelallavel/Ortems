"""Every table importable/exportable with all its columns: round trips, deletes, child tables, workbooks."""

import io

from conftest import Api


def _xlsx(rows: list[list], sheet: str = "Sheet1") -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _sheet(data: bytes, name: str) -> list[list]:
    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO(data))[name]
    return [list(r) for r in ws.iter_rows(values_only=True)]


def _import(api, entity, filename, data, plant, options=None, commit=True):
    job = api.ok(api.post("/imports", data={"entity": entity, "plant_id": plant["id"]}, files={"file": (filename, data)}), 201)
    if options:
        api.ok(api.put(f"/imports/{job['id']}/mapping", json={"mapping": job["mapping"], "options": options}))
    v = api.ok(api.post(f"/imports/{job['id']}/validate"))
    if not commit:
        return v
    assert not v["errors"] or (options or {}).get("skip_invalid_rows"), v["errors"][:5]
    return v, api.ok(api.post(f"/imports/{job['id']}/commit"))


def test_every_registered_table_has_a_template(planner):
    tables = planner.ok(planner.get("/imports/tables"))
    names = {t["table"] for t in tables}
    for must in ("items", "resources", "calendars.shifts", "boms.lines", "production-orders.operations", "setup-matrices.entries", "maintenance", "material-lots", "order-operations", "actual-production"):
        assert must in names, must
    t = next(t for t in tables if t["table"] == "items")
    assert t["key"] == ["code"] and t["fields"][0]["name"] == "id" and t["fields"][-1]["name"] == "delete"


def test_table_roundtrip_update_create_delete(planner, sevilla):
    r = planner.get("/exports/table:customers", params={"format": "xlsx", "plant_id": sevilla["id"]})
    assert r.status_code == 200
    rows = _sheet(r.content, "customers")
    header, first = rows[0], rows[1]
    assert header[0] == "id" and "code" in header and "name" in header
    ni = header.index("name")
    first[ni] = "Renamed in Excel"
    new = [None] * len(header)
    new[header.index("code")] = "C-XL-NEW"
    new[ni] = "Created in Excel"
    v, c = _import(planner, "table:customers", "customers.xlsx", _xlsx([header, first, new]), sevilla)
    assert not v["errors"], v["errors"]
    assert v["stats"]["to_update"] == 1 and v["stats"]["to_create"] == 1
    assert c["status"] == "IMPORTED" and c["stats"]["created"] == 1 and c["stats"]["updated"] == 1
    lst = planner.ok(planner.get("/master-data/customers", params={"q": "C-XL-NEW"}))
    assert lst["items"][0]["name"] == "Created in Excel"
    got = planner.ok(planner.get(f"/master-data/customers/{first[0]}"))
    assert got["name"] == "Renamed in Excel"
    # delete through the sheet
    data = b"code,delete\nC-XL-NEW,yes\n"
    v, c = _import(planner, "table:customers", "del.csv", data, sevilla)
    assert c["stats"]["deleted"] == 1
    assert planner.ok(planner.get("/master-data/customers", params={"q": "C-XL-NEW"}))["total"] == 0


def test_validation_is_a_real_dry_run(planner, sevilla):
    # the item hook rejects max_lot < min_lot: validation must report it and write nothing
    data = b"code,name,min_lot,max_lot\nDRY-1,Dry run,10,5\nDRY-2,Fine,1,5\n"
    v = _import(planner, "table:items", "i.csv", data, sevilla, commit=False)
    assert v["status"] == "INVALID"
    assert any("Maximum lot" in e["message"] and e["row"] == 2 for e in v["errors"])
    assert planner.ok(planner.get("/master-data/items", params={"q": "DRY-"}))["total"] == 0
    assert planner.post(f"/imports/{v['id']}/commit").status_code == 422
    # explicit: import only the valid rows
    planner.ok(planner.put(f"/imports/{v['id']}/mapping", json={"mapping": v["mapping"], "options": {"skip_invalid_rows": True}}))
    c = planner.ok(planner.post(f"/imports/{v['id']}/commit"))
    assert c["stats"]["created"] == 1
    assert planner.ok(planner.get("/master-data/items", params={"q": "DRY-"}))["total"] == 1


def test_child_table_by_parent_code(planner, sevilla):
    cals = planner.ok(planner.get("/master-data/calendars"))["items"]
    code = cals[0]["code"]
    data = f"calendar,weekday,shift_code,start_time,end_time\n{code},6,SUN-X,08:00,12:00\n".encode()
    v, c = _import(planner, "table:calendars.shifts", "shifts.csv", data, sevilla)
    assert not v["errors"], v["errors"]
    assert c["stats"]["created"] == 1
    cal = planner.ok(planner.get(f"/master-data/calendars/{cals[0]['id']}"))
    assert any(sh["shift_code"] == "SUN-X" for sh in cal["shifts"])


def test_unknown_reference_reported(planner, sevilla):
    data = b"code,name,family\nREF-1,x,NO-SUCH-FAMILY\n"
    v = _import(planner, "table:items", "i.csv", data, sevilla, commit=False)
    assert any("unknown product family" in e["message"] for e in v["errors"])


def test_workbook_roundtrip_whole_dataset(client, sevilla):
    admin = Api(client, "admin")
    r = admin.get("/exports/workbook", params={"plant_id": sevilla["id"]})
    assert r.status_code == 200
    from openpyxl import load_workbook

    names = load_workbook(io.BytesIO(r.content), read_only=True).sheetnames
    assert "items" in names and "production-orders" in names and "How to use" in names
    up = admin.ok(admin.post("/imports/workbook", data={"plant_id": sevilla["id"]}, files={"file": ("all.xlsx", r.content)}), 201)
    assert "How to use" in up["ignored_sheets"] and "items" not in up["empty_sheets"]
    ids = [j["id"] for j in up["jobs"]]
    v = admin.ok(admin.post("/imports/batch/validate", json={"job_ids": ids}))
    errs = [(j["entity"], e) for j in v["jobs"] for e in j["errors"]]
    assert not errs, errs[:10]
    # re-importing an unchanged export changes nothing (dates keep their exact local time, references resolve)
    changed = [(j["entity"], j["stats"]["to_create"], j["stats"]["to_update"], j["stats"]["to_delete"]) for j in v["jobs"] if j["stats"]["to_create"] or j["stats"]["to_update"] or j["stats"]["to_delete"]]
    assert not changed, changed
    c = admin.ok(admin.post("/imports/batch/commit", json={"job_ids": ids}))
    assert all(j["status"] == "IMPORTED" for j in c["jobs"])


def test_workbook_new_rows_reference_each_other(client, sevilla):
    admin = Api(client, "admin")
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "product-families"
    ws.append(["code", "name"])
    ws.append(["FAM-WB", "Family from workbook"])
    ws2 = wb.create_sheet("items")
    ws2.append(["code", "name", "family", "item_type", "make_or_buy"])
    ws2.append(["WB-ITEM", "Item from workbook", "FAM-WB", "FINISHED", "MAKE"])
    buf = io.BytesIO()
    wb.save(buf)
    up = admin.ok(admin.post("/imports/workbook", data={"plant_id": sevilla["id"]}, files={"file": ("new.xlsx", buf.getvalue())}), 201)
    ids = [j["id"] for j in up["jobs"]]
    v = admin.ok(admin.post("/imports/batch/validate", json={"job_ids": ids}))
    assert v["can_import"], [j["errors"] for j in v["jobs"]]
    admin.ok(admin.post("/imports/batch/commit", json={"job_ids": ids}))
    it = admin.ok(admin.get("/master-data/items", params={"q": "WB-ITEM"}))["items"][0]
    assert it["family_label"] == "FAM-WB"
