"""API integration tests: security, tenancy, locking, the planning loop, imports/exports, events."""

import io
import uuid

from conftest import PASSWORD, Api


# ------------------------------------------------------------------ security
def test_health(client):
    assert client.get("/health").json()["status"] == "ok"
    assert client.get("/readiness").status_code == 200
    assert "monxu_http_requests_total" in client.get("/metrics").text or client.get("/metrics").status_code == 200


def test_login_failure_and_me(client):
    from fastapi.testclient import TestClient

    c = TestClient(client.app)
    r = c.post("/api/v1/auth/login", json={"username": "planner", "password": "wrong-password"})
    assert r.status_code == 401
    assert c.get("/api/v1/auth/me").status_code == 401
    r = c.post("/api/v1/auth/login", json={"username": "planner", "password": PASSWORD})
    assert r.status_code == 200 and r.json()["csrf_token"]
    assert c.get("/api/v1/auth/me").json()["username"] == "planner"


def test_csrf_required_for_cookie_sessions(planner, sevilla):
    r = planner.c.post("/api/v1/planning/run", json={"scenario_id": sevilla["live_scenario_id"]})  # no header
    assert r.status_code == 403 and r.json()["error"]["code"] == "CSRF"


def test_bearer_token_needs_no_csrf(client, sevilla):
    tok = client.post("/api/v1/auth/token", json={"username": "viewer", "password": PASSWORD}).json()["access_token"]
    from fastapi.testclient import TestClient

    r = TestClient(client.app).get("/api/v1/plants", headers={"Authorization": f"Bearer {tok}"})
    assert r.status_code == 200


def test_rbac_viewer_cannot_plan(client, sevilla):
    viewer = Api(client, "viewer")
    r = viewer.post("/planning/run", json={"scenario_id": sevilla["live_scenario_id"]})
    assert r.status_code == 403 and r.json()["error"]["code"] == "NO_PERMISSION"
    operator = Api(client, "operator")
    assert operator.get("/master-data/items").status_code == 403


def test_tenant_isolation(client, planner, sevilla):
    from monxuplan.services.tenants import create_tenant

    create_tenant("Other Co", "other-co", "otheradmin", "admin@other.example", "Other-Pass-2026", plant_code="OTH")
    other = Api(client, "otheradmin", "Other-Pass-2026", tenant="other-co")
    items = other.ok(other.get("/master-data/items"))
    assert items["total"] == 0
    plants = other.ok(other.get("/plants"))
    assert [p["code"] for p in plants] == ["OTH"]
    # a record of the demo tenant is invisible (404, not 403 — existence is not leaked)
    r = other.get(f"/scenarios/{sevilla['live_scenario_id']}")
    assert r.status_code == 404


def test_optimistic_locking(planner):
    res = planner.ok(planner.get("/master-data/machines", params={"q": "CNC-01"}))["items"][0]
    v = res["version"]
    planner.ok(planner.patch(f"/master-data/machines/{res['id']}", json={"description": "first edit", "version": v}))
    r = planner.patch(f"/master-data/machines/{res['id']}", json={"description": "stale edit", "version": v})
    assert r.status_code == 409 and r.json()["error"]["code"] == "VERSION_CONFLICT"


# ------------------------------------------------------------------ planning loop
def test_plan_run_gantt_explain(planner, base_plan):
    assert base_plan["number"].startswith("PLAN-")
    assert base_plan["hard_violations_placed"] == 0, "the scheduled operations must respect every hard constraint"
    g = planner.ok(planner.get(f"/plans/{base_plan['id']}/gantt"))
    assert g["operations"] and g["resources"]
    op = g["operations"][0]
    detail = planner.ok(planner.get(f"/plans/{base_plan['id']}/operations/{op['id']}"))
    assert detail["scheduled"]["op_key"] == op["id"]
    exp = planner.ok(planner.get(f"/plans/{base_plan['id']}/operations/{op['id']}/explore"))
    assert exp
    bn = planner.ok(planner.get("/bottlenecks", params={"plan_id": base_plan["id"]}))
    assert "ranking_basis" in bn
    k = planner.ok(planner.get("/kpis", params={"plan_id": base_plan["id"]}))
    assert k["values"]["otif"] is not None


def test_move_undo_redo_publish(planner, base_plan, sevilla):
    g = planner.ok(planner.get(f"/plans/{base_plan['id']}/gantt"))
    movable = [o for o in g["operations"] if not o["fixed"] and o["zone"] != "FROZEN"]
    op = movable[len(movable) // 2]
    body = {"op_id": op["id"], "resource_id": op["resource_id"], "start": op["start"], "replan": "DOWNSTREAM", "reason": "test"}
    pv = planner.ok(planner.post(f"/plans/{base_plan['id']}/moves/preview", json=body))
    assert "feasible" in pv and "comparison" in pv
    new = planner.ok(planner.post(f"/plans/{base_plan['id']}/moves", json={**body, "accept_violations": True}))
    assert new["version_no"] > base_plan["version_no"] and new["kind"] == "MANUAL_EDIT"
    back = planner.ok(planner.post(f"/scenarios/{sevilla['live_scenario_id']}/undo"))
    assert back["id"] == base_plan["id"]
    fwd = planner.ok(planner.post(f"/scenarios/{sevilla['live_scenario_id']}/redo"))
    assert fwd["id"] == new["id"]
    pub = planner.ok(planner.post(f"/plans/{new['id']}/publish", json={"reason": "test publish", "force": True}))
    assert pub["status"] == "PUBLISHED"
    d = planner.ok(planner.get("/dispatch", params={"plant_id": sevilla["id"], "hours": 72}))
    assert d["plan"]["id"] == new["id"]


def test_exports(planner, base_plan):
    r = planner.get(f"/plans/{base_plan['id']}/export", params={"format": "xlsx"})
    assert r.status_code == 200
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Schedule", "Orders", "Operations", "Capacity", "Materials", "Alerts", "KPIs"]
    assert wb["Schedule"].max_row > 10
    csv = planner.get(f"/plans/{base_plan['id']}/export", params={"format": "csv", "sheet": "Orders"})
    assert csv.status_code == 200 and "Order" in csv.text.splitlines()[0]


def test_assistant_grounded(planner, base_plan, sevilla):
    a = planner.ok(planner.post("/assistant/ask", json={"question": "What is the bottleneck?", "plant_id": sevilla["id"]}))
    assert a["mode"] == "GROUNDED" and a["intent"] == "BOTTLENECKS" and a["plan"]["id"]
    late = planner.ok(planner.get(f"/plans/{base_plan['id']}"))
    a = planner.ok(planner.post("/assistant/ask", json={"question": "¿Qué pasa si CNC-03 se avería 8 horas?", "plant_id": sevilla["id"]}))
    assert a["intent"] == "WHAT_IF" and a["actions"][0]["kind"] == "BREAKDOWN"
    assert late


def test_dashboard(planner, sevilla, base_plan):
    d = planner.ok(planner.get("/dashboard", params={"plant_id": sevilla["id"]}))
    assert d["tiles"] and d["greeting"]["text"]


# ------------------------------------------------------------------ imports
def test_import_validation_blocks_errors(planner, sevilla):
    data = "code;name;type;safety_stock\nX-1;Good;RAW;5\nX-2;;RAW;abc\nX-1;Duplicate;RAW;1\n"
    job = planner.ok(planner.post("/imports", data={"entity": "items", "plant_id": sevilla["id"]}, files={"file": ("bad.csv", data, "text/csv")}), 201)
    v = planner.ok(planner.post(f"/imports/{job['id']}/validate"))
    assert v["status"] == "INVALID" and not v["stats"]["can_import"]
    msgs = " ".join(e["message"] for e in v["errors"])
    assert "required" in msgs and "not a number" in msgs and "duplicated" in msgs
    r = planner.post(f"/imports/{job['id']}/commit")
    assert r.status_code == 422
    # explicit choice: import only the valid rows
    planner.ok(planner.put(f"/imports/{job['id']}/mapping", json={"mapping": v["mapping"], "options": {"skip_invalid_rows": True}}))
    c = planner.ok(planner.post(f"/imports/{job['id']}/commit"))
    assert c["status"] == "IMPORTED" and c["stats"]["created"] == 1


def test_import_unknown_reference(planner, sevilla):
    data = "number,item_code,quantity,due_date\nT-1,NO-SUCH-ITEM,5,2026-12-01\n"
    job = planner.ok(planner.post("/imports", data={"entity": "production-orders", "plant_id": sevilla["id"]}, files={"file": ("o.csv", data, "text/csv")}), 201)
    v = planner.ok(planner.post(f"/imports/{job['id']}/validate"))
    assert any("unknown item" in e["message"] for e in v["errors"])


def test_export_import_roundtrip(planner, sevilla):
    r = planner.get("/exports/customers", params={"format": "csv"})
    assert r.status_code == 200
    job = planner.ok(planner.post("/imports", data={"entity": "customers", "plant_id": sevilla["id"]}, files={"file": ("customers.csv", r.content, "text/csv")}), 201)
    v = planner.ok(planner.post(f"/imports/{job['id']}/validate"))
    assert v["stats"]["to_create"] == 0 and v["stats"]["to_update"] > 0 and not v["errors"]


# ------------------------------------------------------------------ events
def test_machine_down_event(client, sevilla):
    tok = client.post("/api/v1/auth/token", json={"username": "supervisor", "password": PASSWORD}).json()["access_token"]
    h = {"Authorization": f"Bearer {tok}"}
    cid = uuid.uuid4().hex
    r = client.post("/api/v1/events", json={"type": "MachineDown", "payload": {"resource": "PACK-02", "reason": "sensor fault"}, "correlation_id": cid}, headers=h)
    assert r.status_code == 202, r.text
    again = client.post("/api/v1/events", json={"type": "MachineDown", "payload": {"resource": "PACK-02"}, "correlation_id": cid}, headers=h)
    assert again.json()["duplicate"] is True
    alerts = client.get("/api/v1/alerts", params={"plant_id": sevilla["id"]}, headers=h).json()
    assert any(a["type"] == "MACHINE_BREAKDOWN" and "PACK-02" in a["title"] for a in alerts)
    r = client.post("/api/v1/events", json={"type": "MachineAvailable", "payload": {"resource": "PACK-02"}}, headers=h)
    assert r.status_code == 202


def test_audit_log(client):
    admin = Api(client, "admin")
    a = admin.ok(admin.get("/audit", params={"action": "LOGIN"}))
    assert a["total"] >= 1
